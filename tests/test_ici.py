import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

TOOLKIT_DIR = Path(__file__).resolve().parents[1] / "toolskit"
if str(TOOLKIT_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLKIT_DIR))

import ici_data
import ici_pdf
from ici_excel import COLS, build_workbook, column_span, signature_blocks, date_in_words, reference_group, write_reports

MONTH = "202608"


def tformdet(farmacia, code, lote, precio=1.0, **values):
    """Fila de TFORMDET que cuadra: TOTAL y STOCK_FIN salen de las demás cifras."""
    rec = {"ANNOMES": MONTH, "CODIGO_PRE": farmacia, "CODIGO_MED": code, "MEDLOTE": lote,
           "MEDREGSAN": "RS1", "FEC_EXP": date(2027, 12, 31), "TIPSUM2": "CI", "FFINAN": "DYT",
           "PRECIO": precio, **values}
    salidas = [*ici_data.CONSUMO, "DEVOL", "VENCIDO", "MERMA", "SAL_REGULA", "SAL_CONINS", *ici_data.OTHER_OUT]
    rec["TOTAL"] = sum(rec.get(n, 0) for n in salidas)
    rec["STOCK_FIN"] = sum(rec.get(n, 0) for n in ("SALDO", "INGRE", "REINGRE", "ING_REGULA")) - rec["TOTAL"]
    return rec


def head(tipo, numero, concepto, origen, destino, dia=10, doc=("04", "001-1"), situa="1", refe=""):
    return {"MOVCODITIP": tipo, "MOVNUMERO": numero, "CCTCODIGO": concepto, "ALMORGVIR": origen,
            "ALMDSTVIR": destino, "MOVTIPODCO": doc[0], "MOVNUMEDCO": doc[1], "MOVSITUA": situa,
            "MOVFECHREG": date(2026, 8, dia), "MOVREFE": refe}


def detail(tipo, numero, code, lote, qty, total=None):
    return {"MOVCODITIP": tipo, "MOVNUMERO": numero, "MEDCOD": code, "MEDLOTE": lote, "MOVCANTID": qty,
            "MOVTOTAL": qty if total is None else total}


def month(tformdet_rows, heads=None, details=None, tformato=None):
    tables = {
        "TFORMDET": tformdet_rows,
        "TFORMATO": tformato or [],
        "MALMACEN": [{"ALMCOD": "06502", "ALMDES": "HOSPITAL"}, {"ALMCOD": "06502F01", "ALMDES": "FARMACIA"},
                     {"ALMCOD": "06502F02", "ALMDES": "EMERGENCIA"}, {"ALMCOD": "030A00", "ALMDES": "ALMACEN UNGET"}],
        "MPRODUCTO": [{"MEDCOD": "001", "MEDNOM": "PARACETAMOL", "MEDPRES": "500 mg", "MEDFF": "TABLET"},
                      {"MEDCOD": "002", "MEDNOM": "AMOXICILINA", "MEDPRES": "500 mg", "MEDFF": "TABLET"}],
    }
    if heads is not None:
        tables["TMOVIM"] = heads
        tables["TMOVIMDET"] = details or []
    return ici_data.load_month(lambda name: tables.get(name), MONTH, log=lambda _m: None)


class IciDataTests(unittest.TestCase):
    def test_otras_salidas_keep_the_sismed_total(self):
        rec = tformdet("06502F01", "001", "L1", SALDO=100, VENTA=10, OTRAS_SAL=5, DISTRI=2, SAL_REGULA=3)
        row = ici_data.normalize_row(rec, {})
        self.assertEqual(row["OTRAS"], 7)
        self.assertEqual(ici_data.balance_errors([row]), 0)
        # Si el ajuste ya viene dentro de OTRAS_SAL (no aparte en el total), no se cuenta dos veces.
        inside = {**rec, "TOTAL": rec["TOTAL"] - 3, "STOCK_FIN": rec["STOCK_FIN"] + 3}
        row = ici_data.normalize_row(inside, {})
        self.assertEqual(row["OTRAS"], 4)
        self.assertEqual(ici_data.balance_errors([row]), 0)

    def test_groups_pharmacies_by_establishment_with_recetas(self):
        m = month(
            [tformdet("06502F01", "001", "L1", SALDO=10, SIS=2), tformdet("06502F02", "001", "L1", SALDO=5, VENTA=1),
             tformdet("06502F01", "999", "X", ANNOMES="202607", SALDO=1)],
            tformato=[{"ANNOMES": MONTH, "CODIGO_PRE": "06502F01", "TIPSUM": "C", "REC_SIS": 7, "REC_OTRCON": 2},
                      {"ANNOMES": MONTH, "CODIGO_PRE": "06502F01", "TIPSUM": "D", "REC_SIS": 1},
                      {"ANNOMES": MONTH, "CODIGO_PRE": "06502F02", "TIPSUM": "C", "REC_VTAS": 4}],
        )
        self.assertEqual(len(m.establishments), 1)
        est = m.establishments[0]
        self.assertEqual((est.code, est.name), ("06502", "HOSPITAL"))
        self.assertEqual([p.code for p in est.pharmacies], ["06502F01", "06502F02"])
        f01 = est.pharmacies[0]
        self.assertEqual(f01.recetas[ici_data.CONSUMO.index("SIS")], 8)
        self.assertEqual(f01.recetas[ici_data.CONSUMO.index("OTR_CONV")], 2)
        self.assertEqual(f01.rows[0]["DESC"], "PARACETAMOL 500 mg")
        self.assertEqual(f01.rows[0]["MEDFF"], "TABLET")
        report = ici_data.consolidated_report(m, est)
        self.assertEqual(report.recetas[ici_data.CONSUMO.index("VENTA")], 4)
        self.assertEqual(report.ficha["Farmacias"], "F01 y F02 (consolidado)")

    def test_concept_35_goes_to_vencidos(self):
        m = month(
            [tformdet("06502F01", "001", "L1", SALDO=50, DEVOL=30)],
            heads=[head("S", "1", "35", "06502F0101", "030A0001"), head("S", "2", "07", "06502F0101", "030A0001")],
            details=[detail("S", "1", "001", "L1", 12), detail("S", "2", "001", "L1", 18)],
        )
        row = m.establishments[0].pharmacies[0].rows[0]
        self.assertEqual((row["DEVOL"], row["VENCIDO"]), (18, 12))
        self.assertEqual(ici_data.balance_errors([row]), 0)

    def test_concept_35_already_in_vencidos_is_not_moved_twice(self):
        m = month(
            [tformdet("06502F01", "001", "L1", SALDO=50, DEVOL=5, VENCIDO=12)],
            heads=[head("S", "1", "35", "06502F0101", "030A0001")],
            details=[detail("S", "1", "001", "L1", 12)],
        )
        row = m.establishments[0].pharmacies[0].rows[0]
        self.assertEqual((row["DEVOL"], row["VENCIDO"]), (5, 12))

    def test_cancelled_and_other_month_movements_are_ignored(self):
        m = month(
            [tformdet("06502F01", "001", "L1", SALDO=50, DEVOL=30)],
            heads=[head("S", "1", "35", "06502F0101", "030A0001", situa="0"),
                   {**head("S", "2", "35", "06502F0101", "030A0001"), "MOVFECHREG": date(2026, 9, 1)}],
            details=[detail("S", "1", "001", "L1", 10), detail("S", "2", "001", "L1", 10)],
        )
        self.assertEqual(m.establishments[0].pharmacies[0].rows[0]["VENCIDO"], 0)

    def test_consolidated_nets_deliveries_between_pharmacies(self):
        rows = [
            tformdet("06502F01", "001", "L1", SALDO=100, INGRE=50, DEVOL=40, SIS=10),
            tformdet("06502F02", "001", "L1", REINGRE=30, VENTA=20),
        ]
        heads = [
            head("E", "9", "04", "030A0001", "06502F0101", dia=2, doc=("03", "G-100")),
            head("S", "10", "07", "06502F0101", "06502F0201", dia=5),
            head("E", "11", "07", "06502F0101", "06502F0201", dia=5),
            head("S", "12", "07", "06502F0101", "06502F0201", dia=31),   # llega en septiembre
        ]
        details = [detail("E", "9", "001", "L1", 50), detail("S", "10", "001", "L1", 30),
                   detail("E", "11", "001", "L1", 30), detail("S", "12", "001", "L1", 10)]
        m = month(rows, heads, details)
        est = m.establishments[0]
        report = ici_data.consolidated_report(m, est)
        self.assertEqual(len(report.rows), 1)
        row = report.rows[0]
        self.assertEqual((row["DEVOL"], row["REINGRE"]), (10, 0))
        self.assertEqual(row["STOCK_FIN"], 100 + 50 - 40 - 10 + 30 - 20)
        self.assertEqual(ici_data.balance_errors(report.rows), 0)
        self.assertAlmostEqual(report.internal_value, 30)
        self.assertAlmostEqual(report.in_transit, 10)
        self.assertEqual([g.doc_numero for g in report.guias], ["G-100"])
        # En la farmacia que recibe, la entrega de la F01 sí es una guía de ingreso.
        f02 = ici_data.pharmacy_report(m, est, est.pharmacies[1])
        self.assertEqual([g.numero for g in f02.guias], ["11"])
        self.assertEqual(f02.rows[0]["REINGRE"], 30)

    def test_without_tmovim_the_internal_delivery_is_estimated(self):
        m = month([tformdet("06502F01", "001", "L1", SALDO=100, DEVOL=40),
                   tformdet("06502F02", "001", "L1", REINGRE=40, VENTA=5)])
        report = ici_data.consolidated_report(m, m.establishments[0])
        self.assertTrue(report.internal_estimated)
        self.assertEqual((report.rows[0]["DEVOL"], report.rows[0]["REINGRE"]), (0, 0))
        self.assertEqual(ici_data.balance_errors(report.rows), 0)
        self.assertTrue(m.warnings)

    def test_lot_correction_stays_and_is_marked(self):
        rows = [tformdet("06502F01", "001", "L1", SALDO=20, DEVOL=20),
                tformdet("06502F01", "001", "L2", REINGRE=20)]
        heads = [head("S", "1", "07", "06502F0101", "06502F0101"), head("E", "2", "07", "06502F0101", "06502F0101")]
        details = [detail("S", "1", "001", "L1", 20), detail("E", "2", "001", "L2", 20)]
        m = month(rows, heads, details)
        est = m.establishments[0]
        for report in (ici_data.pharmacy_report(m, est, est.pharmacies[0]), ici_data.consolidated_report(m, est)):
            by_lot = {r["MEDLOTE"]: r for r in report.rows}
            self.assertEqual(by_lot["L1"]["DEVOL"], 20)
            self.assertEqual(by_lot["L2"]["REINGRE"], 20)
            self.assertIn("correccion", by_lot["L1"]["flags"])
            self.assertEqual(len(report.corrections), 2)
            self.assertEqual(report.guias, [])

    def test_month_without_tformdet_is_empty(self):
        m = month([tformdet("06502F01", "001", "L1", ANNOMES="202607", SALDO=1)])
        self.assertEqual(m.establishments, [])


class IciExcelTests(unittest.TestCase):
    def _report(self, **extra):
        rows = [tformdet("06502F01", "001", "L1", SALDO=100, SIS=10, DEVOL=5, **extra),
                tformdet("06502F01", "002", "L9", SALDO=8, FEC_EXP=date(2026, 7, 1))]
        m = month(rows, heads=[head("E", "1", "04", "030A0001", "06502F0101", doc=("03", "G-1"), refe="Intervención sanitaria"),
                               head("E", "2", "04", "030A0001", "06502F0101", doc=("03", "G-2"), refe="SISMED SET")],
                  details=[detail("E", "1", "001", "L1", 0, 0), detail("E", "2", "001", "L1", 0, 12.5)])
        est = m.establishments[0]
        return ici_data.pharmacy_report(m, est, est.pharmacies[0])

    def test_workbook_has_the_approved_sheets(self):
        wb, summary = build_workbook(self._report())
        self.assertEqual(wb.sheetnames, ["ICI", "Vencimientos", "Alertas de farmacia", "Lotes sin movimiento",
                                         "Guías de ingreso", "Correcciones de lote"])
        ws = wb["ICI"]
        headers = [ws.cell(12, c).value for c in range(1, ws.max_column + 1)]
        self.assertNotIn("AJUSTE DE INVENTARIO (+)", headers)
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(ws.print_title_rows, "$11:$12")
        self.assertEqual(summary["sin_movimiento"], 1)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ici.xlsx"
            wb.save(path)
            guides = openpyxl.load_workbook(path)["Guías de ingreso"]
            cells = [c.value for row in guides.iter_rows() for c in row if c.value is not None]
        # Una sección por referencia, en el orden SISMED, intervención, transferencias.
        self.assertLess(cells.index("SISMED  ·  1 guías"), cells.index("INTERVENCIÓN  ·  1 guías"))
        self.assertLess(cells.index("INTERVENCIÓN  ·  1 guías"), cells.index("TRANSFERENCIAS  ·  0 guías"))
        self.assertIn("Distribución", cells)
        self.assertNotIn("04 Distribución", cells)
        self.assertIn("Intervención sanitaria", cells)

    def test_reference_groups_and_date_in_words(self):
        self.assertEqual(reference_group("sismed"), "SISMED")
        self.assertEqual(reference_group("INTERVENCION SANITARIA"), "INTERVENCIÓN")
        self.assertEqual(reference_group("Transf. Hosp. Tarapoto"), "TRANSFERENCIAS")
        self.assertEqual(reference_group(""), "OTRAS REFERENCIAS")
        self.assertEqual(date_in_words(date(2026, 10, 18)), "18 de octubre de 2026")

    def test_signatures_are_centered_on_the_width(self):
        widths = [10] * 20
        self.assertEqual(column_span(widths, 0.10, 0.40), (3, 8))
        self.assertEqual(signature_blocks(widths), ((3, 8), (13, 18)))
        # En el ICI la descripción es ancha: la firma de la derecha es el espejo de la izquierda.
        widths = [c[1] for c in COLS]
        (l1, l2), (r1, r2) = signature_blocks(widths)
        total = sum(widths)
        left_margin = sum(widths[: l1 - 1]) / total
        right_margin = sum(widths[r2:]) / total
        self.assertLess(abs(left_margin - right_margin), 0.02)

    def test_adjustment_column_only_when_used(self):
        wb, _ = build_workbook(self._report(ING_REGULA=4))
        ws = wb["ICI"]
        headers = [ws.cell(12, c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(headers[headers.index("REINGRESOS") + 1], "AJUSTE DE INVENTARIO (+)")


class IciFilesTests(unittest.TestCase):
    def test_one_book_per_pharmacy_and_the_consolidated(self):
        m = month([tformdet("06502F01", "001", "L1", SALDO=10, SIS=2), tformdet("06502F02", "001", "L1", SALDO=5)])
        with tempfile.TemporaryDirectory() as tmp:
            paths = write_reports(m, tmp, pdf=False, log=lambda _m: None)
            names = sorted(Path(p).name for p in paths)
            self.assertEqual(names, ["ICI AGOSTO 2026 - 06502F01 - FARMACIA.xlsx", "ICI AGOSTO 2026 - 06502F02 - EMERGENCIA.xlsx",
                                     "ICI CONSOLIDADO AGOSTO 2026 - 06502 - HOSPITAL.xlsx"])
            only = write_reports(m, tmp, per_pharmacy=False, pdf=False, log=lambda _m: None)
            self.assertEqual([Path(p).name for p in only], ["ICI CONSOLIDADO AGOSTO 2026 - 06502 - HOSPITAL.xlsx"])

    def test_each_book_also_goes_to_pdf(self):
        m = month([tformdet("06502F01", "001", "L1", SALDO=10, SIS=2), tformdet("06502F02", "001", "L1", SALDO=5)])

        class FakeExporter:
            name = "Prueba"
            closed = False
            exported = []

            def export(self, xlsx, pdf):
                if "EMERGENCIA" in xlsx:
                    raise RuntimeError("falló")
                self.exported.append(Path(pdf).name)
                Path(pdf).write_bytes(b"%PDF")

            def close(self):
                FakeExporter.closed = True

        logs = []
        with tempfile.TemporaryDirectory() as tmp:
            paths = write_reports(m, tmp, log=logs.append, open_exporter=lambda log: FakeExporter())
            names = sorted(Path(p).name for p in paths)
        self.assertIn("ICI CONSOLIDADO AGOSTO 2026 - 06502 - HOSPITAL.pdf", names)
        self.assertIn("ICI AGOSTO 2026 - 06502F01 - FARMACIA.pdf", names)
        # Si un PDF falla, el Excel queda y se sigue con los demás.
        self.assertIn("ICI AGOSTO 2026 - 06502F02 - EMERGENCIA.xlsx", names)
        self.assertNotIn("ICI AGOSTO 2026 - 06502F02 - EMERGENCIA.pdf", names)
        self.assertTrue(any("No se pudo pasar a PDF" in line for line in logs))
        self.assertTrue(FakeExporter.closed)

    def test_without_a_spreadsheet_program_only_the_excel_is_made(self):
        logs = []
        exporter = ici_pdf.open_exporter(log=logs.append, com_programs=(), libreoffice=lambda: None)
        self.assertIsNone(exporter)
        self.assertIn("No se pudo generar el PDF", logs[0])

    @unittest.skipUnless(ici_pdf.find_libreoffice(), "sin LibreOffice")
    def test_libreoffice_prints_every_sheet(self):
        m = month([tformdet("06502F01", "001", "L1", SALDO=10, SIS=2)])
        with tempfile.TemporaryDirectory() as tmp:
            paths = write_reports(m, tmp, log=lambda _m: None, open_exporter=lambda log: ici_pdf.open_exporter(
                log=log, com_programs=()))
            pdf = next(p for p in paths if p.endswith(".pdf"))
            data = Path(pdf).read_bytes()
        self.assertTrue(data.startswith(b"%PDF"))
        # Una página por hoja como mínimo: ICI, Vencimientos, Alertas, Lotes sin movimiento, Guías, Correcciones.
        self.assertGreaterEqual(len(re.findall(rb"/Type\s*/Page(?!s)", data)), 6)

    def test_a_single_pharmacy_has_no_consolidated(self):
        m = month([tformdet("06502F01", "001", "L1", SALDO=10)])
        with tempfile.TemporaryDirectory() as tmp:
            paths = write_reports(m, tmp, per_pharmacy=False, pdf=False, log=lambda _m: None)
            self.assertEqual([Path(p).name for p in paths], ["ICI AGOSTO 2026 - 06502F01 - FARMACIA.xlsx"])


if __name__ == "__main__":
    unittest.main()

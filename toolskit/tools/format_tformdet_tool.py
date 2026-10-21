"""
format_tformdet_tool.py
=======================
Herramienta: Format TFORMDET Excel
Categoría  : Herramientas SISMED

Encapsula toda la lógica de formateo del archivo Excel TFORMDET:
  - Selección del .xlsx
  - Validación de cabecera (columna A1 = "CODIGO_EJE")
  - Transformación de columnas (inserciones, eliminaciones, re-ordenamiento)
  - Aplicación de fórmulas CONCATENATE y FECHA CON NUMERO
  - Guardado del archivo resultante con sufijo de mes
"""

import os
import locale
import threading
import flet as ft
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side, NamedStyle

from tools.base_tool import SismedTool, ToolMetadata
import activity


class FormatTformdetTool(SismedTool):
    """Da formato profesional al archivo Excel exportado desde SISMED TFORMDET."""

    # ─── Metadata ──────────────────────────────────────────────────────────────
    @property
    def metadata(self) -> ToolMetadata:
        return ToolMetadata(
            nombre="Format TFORMDET",
            descripcion="Dar formato a archivo excel de Tformdet",
            icono=ft.icons.FORMAT_SHAPES,
            color=ft.colors.INDIGO_400,
            categoria="Herramientas",
            etiquetas=["tformdet", "format", "excel", "xlsx", "formato", "herramientas"],
        )

    # ─── Vista completa ─────────────────────────────────────────────────────────
    def build_view(self, page: ft.Page, navigate_fn) -> ft.Control:
        """Construye la vista de la herramienta de formato TFORMDET."""

        _alert_dialog = None
        _file_name    = ""

        # ── Controles de UI ──────────────────────────────────────────────────
        archivo_seleccionado = ft.TextField(
            label="Archivo .xlsx seleccionado", read_only=True, expand=True,
            border_color=ft.colors.INDIGO_400, border_radius=8,
            content_padding=15, prefix_icon=ft.icons.TABLE_CHART_ROUNDED,
        )
        status_text = ft.Text(
            value="Selecciona un archivo Excel TFORMDET para comenzar",
            size=14, color=ft.colors.ON_SURFACE_VARIANT
        )
        progress_bar = ft.ProgressBar(
            expand=True, visible=False,
            color=ft.colors.INDIGO_400, bgcolor=ft.colors.WHITE24
        )

        # ── Botones (creados antes de las funciones que los referencian) ─────
        seleccionar_button = ft.ElevatedButton(
            content=ft.Row([
                ft.Icon(ft.icons.FILE_OPEN_ROUNDED, size=18, color=ft.colors.WHITE),
                ft.Text("Seleccionar XLSX", weight=ft.FontWeight.BOLD,
                        color=ft.colors.WHITE)
            ]),
            style=ft.ButtonStyle(bgcolor=ft.colors.BLUE_700,
                                 shape=ft.RoundedRectangleBorder(radius=10)),
            height=45,
            on_click=lambda _: file_picker.pick_files(
                allow_multiple=False, allowed_extensions=["xlsx"]
            ),
        )
        modificar_button = ft.ElevatedButton(
            content=ft.Row([
                ft.Icon(ft.icons.AUTO_FIX_HIGH_ROUNDED, size=18,
                        color=ft.colors.WHITE),
                ft.Text("APLICAR FORMATO", weight=ft.FontWeight.BOLD,
                        color=ft.colors.WHITE, size=13),
            ], alignment=ft.MainAxisAlignment.CENTER, spacing=10),
            style=ft.ButtonStyle(
                bgcolor=ft.colors.GREEN_700,
                shape=ft.RoundedRectangleBorder(radius=10),
            ),
            height=45, width=220, visible=False,
        )

        # ── Lógica principal (formateo Excel en hilo separado) ───────────────
        def _close_alert(e):
            nonlocal _alert_dialog
            if _alert_dialog:
                _alert_dialog.open = False
                archivo_seleccionado.value = ""
                status_text.value = "Selecciona un archivo Excel TFORMDET para comenzar"
                status_text.color = ft.colors.ON_SURFACE_VARIANT
                modificar_button.visible = False
                page.update()

        def _modificar_archivo(e):
            nonlocal _alert_dialog, _file_name

            def _process():
                nonlocal _alert_dialog
                modificar_button.visible   = False
                seleccionar_button.disabled = True
                status_text.value  = f"Espere… modificando archivo {_file_name}"
                status_text.color  = ft.colors.AMBER_400
                progress_bar.visible = True
                progress_bar.value = None  # spinner indeterminado
                page.update()

                archivo = archivo_seleccionado.value
                try:
                    workbook = load_workbook(archivo)
                    sheet    = workbook.active

                    # Validar que sea un TFORMDET real
                    if sheet["A1"].value != "CODIGO_EJE":
                        status_text.value = "El archivo seleccionado NO es un TFORMDET.xlsx válido."
                        status_text.color  = ft.colors.RED_400
                        return

                    # Extraer mes del archivo
                    locale.setlocale(locale.LC_TIME, "Spanish_Spain")
                    fecha         = sheet["D2"].value
                    mes           = int(fecha) % 100
                    fecha_completa = datetime(2000, mes, 1)
                    texto_fecha   = fecha_completa.strftime("%B").upper()

                    # Estilos
                    bold_font = Font(bold=True)
                    borde = Border(
                        left=Side(border_style="thin"),
                        right=Side(border_style="thin"),
                        top=Side(border_style="thin"),
                        bottom=Side(border_style="thin"),
                    )

                    # ── Transformaciones de columnas ─────────────────────────
                    sheet.insert_cols(2, 2)
                    sheet.insert_cols(5, 1)
                    sheet.insert_cols(9, 1)

                    headers = {
                        "B1": "COD. UNICO MAS CONSUMO SIS",
                        "C1": "COD UNICO",
                        "E1": "EESS",
                        "I1": "DESCRIPCION MED",
                    }
                    for cell, value in headers.items():
                        sheet[cell] = value

                    sheet.delete_cols(37, 18)
                    for col in [39, 34, 33, 7, 6, 1]:
                        sheet.delete_cols(col)

                    # Copiar LOTE y REG. SANITARIO
                    sheet.insert_cols(7, 2)
                    for row in sheet.iter_rows(
                        min_row=1, max_row=sheet.max_row, min_col=34, max_col=35
                    ):
                        sheet.cell(row=row[0].row, column=7, value=row[0].value)
                        sheet.cell(row=row[0].row, column=8, value=row[1].value)
                    sheet.delete_cols(34, 2)

                    sheet.insert_cols(9, 4)
                    for row in sheet.iter_rows(
                        min_row=1, max_row=sheet.max_row, min_col=37, max_col=40
                    ):
                        for offset, col in enumerate([9, 10, 11, 12]):
                            sheet.cell(row=row[0].row, column=col,
                                       value=row[offset].value)
                    sheet.delete_cols(37, 4)

                    sheet.insert_cols(10, 1)
                    sheet.cell(row=1, column=10, value="FECHA CON NUMERO")

                    # Negrita + bordes en cabecera
                    for col in range(1, sheet.max_column + 1):
                        if sheet.cell(row=1, column=col).value is not None:
                            sheet.cell(row=1, column=col).font   = bold_font
                            sheet.cell(row=1, column=col).border = borde

                    # Formato de fecha en columna I
                    try:
                        date_style = NamedStyle(name="datetime",
                                               number_format="DD/MM/YYYY")
                    except ValueError:
                        date_style = workbook.named_styles["datetime"]

                    for cell in sheet["I"]:
                        if cell.row == 1:
                            continue
                        try:
                            if cell.value and isinstance(cell.value, str):
                                cell.value = datetime.strptime(
                                    cell.value, "%Y-%m-%d"
                                )
                            cell.style = date_style
                        except Exception:
                            pass

                    # Fórmulas
                    formula_fecha = (
                        "=CONCATENATE(YEAR(I2), "
                        "IF(MONTH(I2)<10, CONCATENATE(0, MONTH(I2)), MONTH(I2)), "
                        "IF(DAY(I2)<10, CONCATENATE(0, DAY(I2)), DAY(I2)))"
                    )
                    formula_cod_unico     = "=CONCATENATE(C2,E2)"
                    formula_cod_mas_cons  = "=CONCATENATE(C2,E2,R2,S2)"

                    last_row = sheet.max_row
                    for row in range(2, last_row + 1):
                        sheet[f"J{row}"].value = formula_fecha.replace("I2", f"I{row}")
                        sheet[f"B{row}"].value = formula_cod_unico
                        sheet[f"A{row}"].value = formula_cod_mas_cons

                    # Guardar
                    nuevo_archivo = archivo.replace(".xlsx", f"-{texto_fecha}.xlsx")
                    workbook.save(nuevo_archivo)

                    status_text.value = "✔ Formato aplicado correctamente."
                    status_text.color  = ft.colors.GREEN_400
                    progress_bar.value = 1
                    page.update()

                    _alert_dialog = ft.AlertDialog(
                        title=ft.Row([
                            ft.Icon(ft.icons.CHECK_CIRCLE_ROUNDED,
                                    color=ft.colors.GREEN_500),
                            ft.Text("¡Completado con Éxito!",
                                    weight=ft.FontWeight.W_900)
                        ]),
                        content=ft.Text(
                            f"El archivo formateado se guardó en:\n\n{nuevo_archivo}",
                            size=13
                        ),
                        actions=[ft.ElevatedButton("Entendido",
                                                   on_click=_close_alert)],
                        open=True,
                        shape=ft.RoundedRectangleBorder(radius=16)
                    )
                    page.overlay.append(_alert_dialog)

                except PermissionError as err:
                    status_text.value = f"Error: Archivo en uso. Ciérralo primero. ({err})"
                    status_text.color  = ft.colors.RED_400
                except Exception as ex:
                    status_text.value = f"Error inesperado: {str(ex)}"
                    status_text.color  = ft.colors.RED_400
                finally:
                    progress_bar.visible    = False
                    seleccionar_button.disabled = False
                    page.update()

            activity.background("Formato TFORMDET", _process)

        modificar_button.on_click = _modificar_archivo

        # ── FilePicker ───────────────────────────────────────────────────────
        def _on_file_picked(e):
            nonlocal _file_name
            if e.files:
                fp = e.files[0].path
                if not fp.lower().endswith(".xlsx"):
                    archivo_seleccionado.value = ""
                    status_text.value = "Por favor, selecciona un archivo con extensión .xlsx."
                    status_text.color = ft.colors.RED_400
                    modificar_button.visible = False
                else:
                    archivo_seleccionado.value = fp
                    _file_name = os.path.basename(fp)
                    status_text.value = f"ARCHIVO A MODIFICAR: {_file_name}"
                    status_text.color = ft.colors.TEAL_300
                    modificar_button.visible = True
            else:
                archivo_seleccionado.value = ""
                status_text.value = "Selecciona un archivo Excel TFORMDET para comenzar"
                status_text.color = ft.colors.ON_SURFACE_VARIANT
                modificar_button.visible = False
            page.update()

        file_picker = ft.FilePicker(on_result=_on_file_picked)
        page.overlay.append(file_picker)
        page.update()

        # ── Layout ──────────────────────────────────────────────────────────
        return ft.Container(
            padding=40,
            expand=True,
            content=ft.Column(
                scroll=ft.ScrollMode.AUTO,
                controls=[
                # Cabecera Premium
                self._section_header(
                    self.metadata.icono,
                    self.metadata.color,
                    "Format TFORMDET Excel",
                    "Aplica formato profesional al Excel exportado de SISMED",
                    right_content=self._back_button("Volver a Herramientas", navigate_fn, 2)
                ),

                ft.Divider(height=10, color=ft.colors.TRANSPARENT),

                # Tarjeta central
                ft.Row([
                    ft.Container(
                        bgcolor=ft.colors.BLUE_GREY_900,
                        padding=ft.padding.only(top=35, bottom=35,
                                                left=40, right=40),
                        border_radius=16,
                        width=800,
                        content=ft.Column([
                            ft.Row([
                                ft.Icon(ft.icons.SETTINGS_ROUNDED,
                                        color=ft.colors.INDIGO_400, size=24),
                                ft.Text("Parámetros de Formato", size=16,
                                        weight=ft.FontWeight.BOLD,
                                        color=ft.colors.WHITE),
                            ], spacing=10),
                            ft.Divider(color=ft.colors.WHITE24, height=15),

                            ft.Text("Archivo TFORMDET .xlsx a procesar:",
                                    size=13, color=ft.colors.GREY_400),
                            ft.Row([archivo_seleccionado, seleccionar_button],
                                   spacing=15,
                                   vertical_alignment=ft.CrossAxisAlignment.CENTER),

                            ft.Container(height=20),
                            status_text,

                            ft.Container(height=40),
                            progress_bar,
                            ft.Container(height=5),

                            ft.Row([modificar_button],
                                   alignment=ft.MainAxisAlignment.END),
                        ])
                    )
                ], alignment=ft.MainAxisAlignment.CENTER),

            ], spacing=10)
        )

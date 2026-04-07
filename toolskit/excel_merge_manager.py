import os
from collections import Counter
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook

try:
    import xlrd
except ImportError:
    xlrd = None


@dataclass
class MergeExcelResult:
    output_path: str
    reference_headers: list[str] = field(default_factory=list)
    included_files: list[str] = field(default_factory=list)
    excluded_files: list[dict] = field(default_factory=list)
    total_rows_written: int = 0


class ExcelMergeManager:
    """Une archivos Excel en uno solo validando compatibilidad de cabeceras."""

    def __init__(self, log_callback=None):
        self._log_fn = log_callback

    def log(self, message: str):
        if self._log_fn:
            self._log_fn(message)
        else:
            print(message)

    @staticmethod
    def _trim_row(values) -> list:
        row = list(values or [])
        while row and (row[-1] is None or str(row[-1]).strip() == ""):
            row.pop()
        return row

    @staticmethod
    def _is_empty_row(values) -> bool:
        return not any(value is not None and str(value).strip() != "" for value in values or [])

    @staticmethod
    def _normalize_header(value) -> str:
        text = "" if value is None else str(value)
        return " ".join(text.strip().split()).upper()

    @staticmethod
    def _get_first_visible_openpyxl_sheet(workbook):
        for sheet in workbook.worksheets:
            if getattr(sheet, "sheet_state", "visible") == "visible":
                return sheet
        return None

    @staticmethod
    def _get_first_visible_xlrd_sheet(workbook):
        for sheet_index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_index)
            if getattr(sheet, "visibility", 0) == 0:
                return sheet
        return None

    def _extract_header_and_rows(self, row_iterable) -> tuple[list[str], list[list]]:
        header = None
        rows = []

        for raw_values in row_iterable:
            raw_row = list(raw_values or [])
            trimmed_values = self._trim_row(raw_row)
            if header is None:
                if self._is_empty_row(trimmed_values):
                    continue
                header = ["" if value is None else str(value).strip() for value in trimmed_values]
                continue

            if self._is_empty_row(trimmed_values):
                continue

            row = raw_row
            if len(row) < len(header):
                row.extend([None] * (len(header) - len(row)))
            rows.append(row[:len(header)])

        if header is None:
            raise ValueError("No se encontró una fila de cabeceras válida en la primera hoja.")

        normalized = [self._normalize_header(value) for value in header]
        if not all(normalized):
            raise ValueError("Se detectaron cabeceras vacías o no válidas en la primera hoja.")

        return header, rows

    def _read_first_sheet_xlsx(self, file_path: str) -> tuple[list[str], list[list]]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = self._get_first_visible_openpyxl_sheet(workbook)
            if sheet is None:
                raise ValueError("El archivo no contiene hojas visibles para procesar.")

            return self._extract_header_and_rows(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()

    def _read_first_sheet_xls(self, file_path: str) -> tuple[list[str], list[list]]:
        if xlrd is None:
            raise ValueError(
                "El archivo está en formato .xls y falta el soporte para leerlo. Instala xlrd o convierte el archivo a .xlsx."
            )

        workbook = xlrd.open_workbook(file_path, on_demand=True)
        try:
            if workbook.nsheets < 1:
                raise ValueError("El archivo no contiene hojas visibles para procesar.")

            sheet = self._get_first_visible_xlrd_sheet(workbook)
            if sheet is None:
                raise ValueError("El archivo no contiene hojas visibles para procesar.")

            def _rows():
                for row_index in range(sheet.nrows):
                    yield sheet.row_values(row_index)

            return self._extract_header_and_rows(_rows())
        finally:
            workbook.release_resources()

    def _read_first_sheet(self, file_path: str) -> tuple[list[str], list[list]]:
        extension = os.path.splitext(file_path)[1].lower()
        if extension == ".xls":
            return self._read_first_sheet_xls(file_path)
        return self._read_first_sheet_xlsx(file_path)

    def _describe_header_mismatch(
        self,
        reference_headers: list[str],
        reference_normalized: list[str],
        current_headers: list[str],
        current_normalized: list[str],
    ) -> str:
        reference_counter = Counter(reference_normalized)
        current_counter = Counter(current_normalized)

        missing = list((reference_counter - current_counter).elements())
        extra = list((current_counter - reference_counter).elements())

        message_parts = []
        if missing:
            message_parts.append("faltan columnas: " + ", ".join(sorted(set(missing))))
        if extra:
            message_parts.append("sobran columnas: " + ", ".join(sorted(set(extra))))

        if not message_parts:
            message_parts.append(
                "las cabeceras tienen el mismo contenido pero no se pudieron reordenar de forma segura"
            )

        return "; ".join(message_parts)

    def merge_files(self, file_paths: list[str], output_path: str) -> MergeExcelResult:
        unique_paths = []
        seen = set()
        for path in file_paths:
            normalized = os.path.normcase(os.path.abspath(path))
            if normalized in seen:
                continue
            seen.add(normalized)
            unique_paths.append(path)

        if len(unique_paths) < 2:
            raise ValueError("Selecciona al menos 2 archivos Excel para unir.")

        output_path = os.path.abspath(output_path)
        output_dir = os.path.dirname(output_path)
        os.makedirs(output_dir, exist_ok=True)

        reference_headers = None
        reference_normalized = None
        combined_rows = []
        included_files = []
        excluded_files = []
        duplicates_in_reference = False

        self.log(f"[INFO] Archivos seleccionados: {len(unique_paths)}")

        for file_path in unique_paths:
            file_name = os.path.basename(file_path)
            self.log(f"[INFO] Validando {file_name}...")
            try:
                headers, rows = self._read_first_sheet(file_path)
                normalized_headers = [self._normalize_header(value) for value in headers]

                if reference_headers is None:
                    reference_headers = headers
                    reference_normalized = normalized_headers
                    duplicates_in_reference = len(set(reference_normalized)) != len(reference_normalized)
                    combined_rows.extend(rows)
                    included_files.append(file_path)
                    self.log(f"  [OK] Archivo base aceptado con {len(rows)} filas de datos.")
                    continue

                if normalized_headers == reference_normalized:
                    combined_rows.extend(rows)
                    included_files.append(file_path)
                    self.log(f"  [OK] Cabeceras compatibles. Se agregaron {len(rows)} filas.")
                    continue

                if not duplicates_in_reference and Counter(normalized_headers) == Counter(reference_normalized):
                    current_index = {name: idx for idx, name in enumerate(normalized_headers)}
                    reordered_rows = []
                    for row in rows:
                        reordered_rows.append([
                            row[current_index[name]] if current_index[name] < len(row) else None
                            for name in reference_normalized
                        ])
                    combined_rows.extend(reordered_rows)
                    included_files.append(file_path)
                    self.log(
                        f"  [OK] Cabeceras compatibles con orden distinto. Se reordenaron columnas y se agregaron {len(rows)} filas."
                    )
                    continue

                reason = self._describe_header_mismatch(
                    reference_headers,
                    reference_normalized,
                    headers,
                    normalized_headers,
                )
                excluded_files.append({"path": file_path, "reason": reason})
                self.log(f"  [WARN] Archivo excluido: {reason}.")
            except Exception as exc:
                reason = str(exc)
                excluded_files.append({"path": file_path, "reason": reason})
                self.log(f"  [WARN] Archivo excluido: {reason}.")

        if len(included_files) < 2:
            raise ValueError(
                "No se pudo generar el consolidado porque se necesitan al menos 2 archivos válidos con cabeceras compatibles."
            )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Unificado"
        sheet.append(reference_headers)
        for row in combined_rows:
            sheet.append(row)
        workbook.save(output_path)
        workbook.close()

        self.log(f"[OK] Archivo consolidado generado: {output_path}")
        self.log(f"[INFO] Resumen: {len(included_files)} aceptados, {len(excluded_files)} excluidos, {len(combined_rows)} filas unificadas.")

        return MergeExcelResult(
            output_path=output_path,
            reference_headers=reference_headers,
            included_files=included_files,
            excluded_files=excluded_files,
            total_rows_written=len(combined_rows),
        )
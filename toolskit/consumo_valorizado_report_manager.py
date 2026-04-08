"""
consumo_valorizado_report_manager.py
====================================
Motor para generar el reporte de consumo valorizado por establecimiento.

Replica la logica de una consulta TRANSFORM/PIVOT de Access:
  TRANSFORM Sum([CAMPO]*[PRECIO]) AS TOTAL
  SELECT ANNOMES, ALMCOD, ALMDES
  GROUP BY ANNOMES, ALMCOD, ALMDES
  PIVOT MEDTIP

Tipos soportados:
  - SIS
  - INTERSAN
  - VENTA
"""

import os
import tempfile
from collections import defaultdict

from dbf_utils import open_dbf, snapshot_dbf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ConsumoValorizadoReportManager:
    """Genera un reporte pivotado por MEDTIP valorizado por establecimiento."""

    _PERIOD_CACHE: dict[str, tuple[tuple[int, int], list[str]]] = {}
    SUMMARY_SHEET_TITLE = "Resumen Total"
    SUMMARY_REPORT_ORDER = ["sis", "intersan", "venta"]

    REPORT_TYPES = {
        "sis": {
            "field": "SIS",
            "label": "SIS",
            "file_label": "SIS",
            "sheet_title": "Valorizado SIS",
            "summary_label": "SIS",
        },
        "intersan": {
            "field": "INTERSAN",
            "label": "INTERVENCION SANITARIA",
            "file_label": "INTERVENCION_SANITARIA",
            "sheet_title": "Valorizado InterSan",
            "summary_label": "INTERV. SANITARIA",
        },
        "venta": {
            "field": "VENTA",
            "label": "VENTA",
            "file_label": "VENTA",
            "sheet_title": "Valorizado Venta",
            "summary_label": "VENTA",
        },
    }

    def __init__(self, sismed_dir: str, log_callback=None):
        self.sismed_dir = sismed_dir.strip()
        self._log_fn = log_callback

    def log(self, msg: str):
        if self._log_fn:
            self._log_fn(msg)
        else:
            print(msg)

    def _dbf(self, nombre: str) -> str:
        return os.path.join(self.sismed_dir, "DATOS", nombre)

    def _read_dbf(self, path: str) -> list:
        if not os.path.exists(path):
            self.log(f"  ⚠ No encontrado: {path}")
            return []
        try:
            with tempfile.TemporaryDirectory(prefix="toolkit_valorizado_") as snapshot_dir:
                read_path = snapshot_dbf(path, snapshot_dir)
                return list(open_dbf(read_path))
        except Exception as exc:
            self.log(f"  ✗ Error leyendo {os.path.basename(path)}: {exc}")
            return []

    @staticmethod
    def _s(val) -> str:
        return str(val).strip() if val is not None else ""

    @staticmethod
    def _f(val) -> float:
        if val is None:
            return 0.0
        if isinstance(val, (int, float)):
            return float(val)

        text = str(val).strip()
        if not text:
            return 0.0

        if "," in text and "." in text:
            text = text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")

        try:
            return float(text)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _medtip_sort_key(value: str):
        text = str(value or "").strip()
        if text.isdigit():
            return (0, int(text), text)
        return (1, text.upper(), text)

    @staticmethod
    def _summary_bucket(medtip: str) -> str:
        text = str(medtip or "").strip().upper()
        compact = text.lstrip("0") or "0"
        if text in {"M", "MEDICAMENTO", "MEDICAMENTOS"} or compact == "1":
            return "medicamentos"
        if text in {"I", "INSUMO", "INSUMOS"} or compact == "2":
            return "insumos"
        return "insumos"

    @classmethod
    def _should_create_total_summary(cls, selected_report_types: list[str]) -> bool:
        return set(selected_report_types) == set(cls.SUMMARY_REPORT_ORDER)

    def _load_medicam(self) -> dict:
        catalog = {}
        for rec in self._read_dbf(self._dbf("MMEDICAM.DBF")):
            cod = self._s(rec.get("MEDCOD")).zfill(5)
            if not cod or cod == "00000":
                continue

            medtip = self._s(rec.get("MEDTIP")) or "SIN_TIPO"
            catalog[cod] = {
                "medtip": medtip,
            }

        self.log(f"  ✓ Medicamentos cargados: {len(catalog)}")
        return catalog

    def _load_almacen(self) -> dict:
        catalog = {}
        for rec in self._read_dbf(self._dbf("MALMACEN.DBF")):
            cod = self._s(rec.get("ALMCOD"))
            if cod:
                catalog[cod] = self._s(rec.get("ALMDES"))

        self.log(f"  ✓ Almacenes cargados: {len(catalog)}")
        return catalog

    @staticmethod
    def _is_valid_period(value: str) -> bool:
        return bool(value) and len(value) == 6 and value.isdigit() and 1 <= int(value[4:6]) <= 12

    @classmethod
    def _build_period_range(cls, start_period: str, end_period: str) -> list[str]:
        if not cls._is_valid_period(start_period) or not cls._is_valid_period(end_period):
            return []

        start_year, start_month = int(start_period[:4]), int(start_period[4:6])
        end_year, end_month = int(end_period[:4]), int(end_period[4:6])
        if (start_year, start_month) > (end_year, end_month):
            return []

        periods = []
        year, month = start_year, start_month
        while (year, month) <= (end_year, end_month):
            periods.append(f"{year}{month:02d}")
            month += 1
            if month > 12:
                month = 1
                year += 1
            if len(periods) > 240:
                break
        return periods

    @staticmethod
    def _extract_annomes_periods(path: str) -> list[str]:
        with open(path, "rb") as dbf_file:
            header = dbf_file.read(32)
            if len(header) < 32:
                return []

            header_length = int.from_bytes(header[8:10], "little")
            record_length = int.from_bytes(header[10:12], "little")
            record_count = int.from_bytes(header[4:8], "little")
            if header_length <= 32 or record_length <= 1:
                return []

            descriptor_bytes = dbf_file.read(header_length - 32)
            field_offset = 1
            annomes_offset = None
            annomes_length = None

            for index in range(0, len(descriptor_bytes), 32):
                descriptor = descriptor_bytes[index:index + 32]
                if not descriptor or descriptor[0] == 0x0D:
                    break

                name = descriptor[:11].split(b"\x00", 1)[0].decode("ascii", errors="ignore").strip()
                length = descriptor[16]

                if name.upper() == "ANNOMES":
                    annomes_offset = field_offset
                    annomes_length = length
                    break

                field_offset += length

            if annomes_offset is None or annomes_length is None:
                return []

            dbf_file.seek(header_length)
            periods = set()
            for _ in range(record_count):
                record = dbf_file.read(record_length)
                if len(record) < record_length:
                    break
                if record[:1] == b"*":
                    continue

                raw_value = record[annomes_offset:annomes_offset + annomes_length]
                value = raw_value.decode("latin-1", errors="ignore").replace("\x00", "").strip()
                if value and len(value) == 6 and value.isdigit() and 1 <= int(value[4:6]) <= 12:
                    periods.add(value)

            return sorted(periods, reverse=True)

    def get_available_periods(self) -> list[str]:
        source_path = self._dbf("TFORMDET.DBF")
        if not os.path.exists(source_path):
            self.log(f"  ⚠ No encontrado: {source_path}")
            return []

        stat = os.stat(source_path)
        cache_key = (stat.st_mtime_ns, stat.st_size)
        cached = self._PERIOD_CACHE.get(source_path)
        if cached and cached[0] == cache_key:
            result = cached[1]
            self.log(f"  ✓ Periodos detectados en TFORMDET: {len(result)}")
            return result

        try:
            with tempfile.TemporaryDirectory(prefix="toolkit_valorizado_periods_") as snapshot_dir:
                read_path = snapshot_dbf(source_path, snapshot_dir)
                result = self._extract_annomes_periods(read_path)
            self._PERIOD_CACHE[source_path] = (cache_key, result)
            self.log(f"  ✓ Periodos detectados en TFORMDET: {len(result)}")
            return result
        except Exception as exc:
            self.log(f"  ✗ Error leyendo periodos de TFORMDET: {exc}")
            return []

    @classmethod
    def _normalize_report_types(cls, report_types) -> tuple[list[str], list[str]]:
        if isinstance(report_types, str):
            raw_types = [report_types]
        else:
            raw_types = list(report_types or [])

        selected: list[str] = []
        invalid: list[str] = []
        for raw_type in raw_types:
            normalized = str(raw_type or "").strip().lower()
            if not normalized:
                continue
            if normalized not in cls.REPORT_TYPES:
                invalid.append(normalized)
                continue
            if normalized not in selected:
                selected.append(normalized)
        return selected, invalid

    def generate(
        self,
        output_path: str,
        report_types,
        period_start: str,
        period_end: str | None = None,
    ) -> bool:
        selected_report_types, invalid_report_types = self._normalize_report_types(report_types)
        if invalid_report_types:
            self.log("✗ Tipos de reporte no soportados: " + ", ".join(invalid_report_types))
            return False
        if not selected_report_types:
            self.log("✗ Debe seleccionar al menos un tipo de reporte.")
            return False
        if not self._is_valid_period(period_start):
            self.log(f"✗ Periodo inicial invalido: {period_start}")
            return False

        period_end = period_end or period_start
        if not self._is_valid_period(period_end):
            self.log(f"✗ Periodo final invalido: {period_end}")
            return False

        selected_periods = self._build_period_range(period_start, period_end)
        if not selected_periods:
            self.log("✗ El periodo inicial no puede ser mayor al periodo final.")
            return False

        selected_periods_set = set(selected_periods)
        period_label = period_start if period_start == period_end else f"{period_start} -> {period_end}"

        self.log("▶ Iniciando reporte de consumo valorizado...")
        if len(selected_report_types) == 1:
            self.log(f"  Tipo     : {self.REPORT_TYPES[selected_report_types[0]]['label']}")
        else:
            self.log(
                "  Tipos    : "
                + ", ".join(self.REPORT_TYPES[report_type]["label"] for report_type in selected_report_types)
            )
        self.log(f"  Periodo  : {period_label}")
        for report_type in selected_report_types:
            report_def = self.REPORT_TYPES[report_type]
            self.log(f"  Formula  : Sum([{report_def['field']}]*[PRECIO]) -> {report_def['sheet_title']}")

        try:
            cat_med = self._load_medicam()
            cat_alm = self._load_almacen()
            tformdet = self._read_dbf(self._dbf("TFORMDET.DBF"))

            if not tformdet:
                self.log("✗ No se encontraron registros en TFORMDET.DBF.")
                return False

            self.log(f"  ✓ Registros TFORMDET: {len(tformdet)}")

            available_periods = {
                annomes
                for rec in tformdet
                if self._is_valid_period(annomes := self._s(rec.get("ANNOMES")))
            }
            missing_periods = [period for period in selected_periods if period not in available_periods]
            if missing_periods:
                self.log(
                    "✗ No se encontraron datos para los siguientes periodos: "
                    + ", ".join(missing_periods)
                )
                return False

            include_total_summary = self._should_create_total_summary(selected_report_types)

            data_by_type = {
                report_type: defaultdict(lambda: defaultdict(float))
                for report_type in selected_report_types
            }
            medtips_by_type = {
                report_type: set()
                for report_type in selected_report_types
            }
            summary_data = defaultdict(
                lambda: {
                    report_type: {"medicamentos": 0.0, "insumos": 0.0}
                    for report_type in self.SUMMARY_REPORT_ORDER
                }
            )
            sheet_stats = {
                report_type: {"rows_exported": 0}
                for report_type in selected_report_types
            }
            stats = {
                "invalid": 0,
                "missing_med": 0,
                "missing_alm": 0,
                "matched": 0,
                "period_rows": 0,
            }

            for rec in tformdet:
                annomes = self._s(rec.get("ANNOMES"))
                if annomes not in selected_periods_set:
                    continue

                stats["period_rows"] += 1
                cod_pre = self._s(rec.get("CODIGO_PRE"))
                cod_med = self._s(rec.get("CODIGO_MED")).zfill(5)

                if (
                    not annomes
                    or len(annomes) != 6
                    or not annomes.isdigit()
                    or not cod_pre
                    or not cod_med
                    or cod_med == "00000"
                ):
                    stats["invalid"] += 1
                    continue

                med_info = cat_med.get(cod_med)
                if med_info is None:
                    stats["missing_med"] += 1
                    continue

                alm_desc = cat_alm.get(cod_pre)
                if alm_desc is None:
                    stats["missing_alm"] += 1
                    continue

                precio = self._f(rec.get("PRECIO"))
                medtip = med_info["medtip"]
                summary_bucket = self._summary_bucket(medtip)

                for report_type in selected_report_types:
                    field_name = self.REPORT_TYPES[report_type]["field"]
                    total_value = self._f(rec.get(field_name)) * precio
                    data_by_type[report_type][(annomes, cod_pre, alm_desc)][medtip] += total_value
                    medtips_by_type[report_type].add(medtip)
                    if include_total_summary:
                        summary_data[annomes][report_type][summary_bucket] += total_value

                stats["matched"] += 1

            if not stats["period_rows"]:
                self.log(f"✗ El rango {period_label} no tiene registros en TFORMDET.")
                return False

            if not any(data_by_type[report_type] for report_type in selected_report_types):
                self.log("✗ No se obtuvieron filas validas tras relacionar TFORMDET, MALMACEN y MMEDICAM.")
                return False

            wb = Workbook()
            wb.remove(wb.active)
            sheets_created = 0

            for report_type in selected_report_types:
                report_def = self.REPORT_TYPES[report_type]
                report_data = data_by_type[report_type]
                if not report_data:
                    self.log(f"  ⚠ {report_def['label']}: sin filas válidas para el periodo seleccionado.")
                    continue

                medtip_headers = sorted(medtips_by_type[report_type], key=self._medtip_sort_key)
                sorted_keys = sorted(report_data.keys(), key=lambda item: (item[0], item[1], item[2].upper()))
                ws = wb.create_sheet(title=report_def["sheet_title"])
                header = ["ANNOMES", "ALMCOD", "ALMDES"] + medtip_headers + ["TOTAL GENERAL"]
                self._write_header(ws, header)

                row_n = 2
                for annomes, cod_pre, alm_desc in sorted_keys:
                    totals_by_medtip = report_data[(annomes, cod_pre, alm_desc)]

                    row = [annomes, cod_pre, alm_desc]
                    total_general = 0.0
                    for medtip in medtip_headers:
                        value = totals_by_medtip.get(medtip, 0.0)
                        row.append(value)
                        total_general += value
                    row.append(total_general)
                    ws.append(row)

                    ws.cell(row=row_n, column=1).number_format = "@"
                    ws.cell(row=row_n, column=2).number_format = "@"
                    ws.cell(row=row_n, column=1).value = annomes
                    ws.cell(row=row_n, column=2).value = cod_pre
                    row_n += 1

                self._auto_style(ws, len(header), row_n - 1)
                sheet_stats[report_type]["rows_exported"] = row_n - 2
                sheets_created += 1

            if include_total_summary:
                summary_ws = wb.create_sheet(title=self.SUMMARY_SHEET_TITLE)
                summary_rows = self._write_total_summary_sheet(summary_ws, selected_periods, summary_data)
                sheets_created += 1
            else:
                summary_rows = 0

            if sheets_created == 0:
                self.log("✗ No se pudieron generar hojas para los reportes seleccionados.")
                return False

            wb.save(output_path)

            self.log(f"  ✓ Periodos procesados: {len(selected_periods)}")
            self.log(f"  ✓ Registros del periodo: {stats['period_rows']}")
            self.log(f"  ✓ Hojas generadas   : {sheets_created}")
            for report_type in selected_report_types:
                report_def = self.REPORT_TYPES[report_type]
                rows_exported = sheet_stats[report_type]["rows_exported"]
                if rows_exported:
                    self.log(f"  ✓ {report_def['sheet_title']}: {rows_exported} filas exportadas")
            if summary_rows:
                self.log(f"  ✓ {self.SUMMARY_SHEET_TITLE}: {summary_rows} filas exportadas")
            if stats["invalid"]:
                self.log(f"  ⚠ Registros omitidos por datos incompletos: {stats['invalid']}")
            if stats["missing_med"]:
                self.log(f"  ⚠ Registros sin MMEDICAM relacionado: {stats['missing_med']}")
            if stats["missing_alm"]:
                self.log(f"  ⚠ Registros sin MALMACEN relacionado: {stats['missing_alm']}")
            self.log(f"✅ Reporte guardado en: {output_path}")
            return True

        except Exception as exc:
            import traceback

            self.log(f"✗ ERROR: {exc}")
            self.log(traceback.format_exc())
            return False

    def _write_header(self, ws, columns: list[str]):
        fill = PatternFill("solid", fgColor="0F4C5C")
        font = Font(bold=True, color="FFFFFF", size=10)
        border = Border(
            left=Side(style="thin", color="2A6877"),
            right=Side(style="thin", color="2A6877"),
            top=Side(style="thin", color="2A6877"),
            bottom=Side(style="medium", color="2A6877"),
        )
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for idx, name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=idx, value=name)
            cell.fill = fill
            cell.font = font
            cell.border = border
            cell.alignment = align

        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    def _write_total_summary_sheet(self, ws, selected_periods: list[str], summary_data: dict) -> int:
        columns = ["N°", "AÑO/MES", "CONSUMO ICI", "MEDICAMENTOS", "INSUMOS", "SUB-TOTAL", "TOTAL"]
        self._write_header(ws, columns)

        row_n = 2
        group_ranges: list[tuple[int, int]] = []
        for index, annomes in enumerate(selected_periods, start=1):
            group_start = row_n
            period_rows = []
            period_total = 0.0
            period_summary = summary_data.get(annomes, {})

            for report_type in self.SUMMARY_REPORT_ORDER:
                totals = period_summary.get(report_type, {})
                medicamentos = float(totals.get("medicamentos", 0.0))
                insumos = float(totals.get("insumos", 0.0))
                subtotal = medicamentos + insumos
                period_total += subtotal
                period_rows.append((
                    self.REPORT_TYPES[report_type]["summary_label"],
                    medicamentos,
                    insumos,
                    subtotal,
                ))

            for summary_label, medicamentos, insumos, subtotal in period_rows:
                ws.cell(row=row_n, column=3, value=summary_label)
                ws.cell(row=row_n, column=4, value=medicamentos)
                ws.cell(row=row_n, column=5, value=insumos)
                ws.cell(row=row_n, column=6, value=subtotal)
                row_n += 1

            group_end = row_n - 1
            ws.cell(row=group_start, column=1, value=f"{index:02d}")
            ws.cell(row=group_start, column=2, value=annomes)
            ws.cell(row=group_start, column=7, value=period_total)
            group_ranges.append((group_start, group_end))

        self._auto_style_total_summary(ws, row_n - 1, group_ranges)
        return row_n - 2

    def _auto_style_total_summary(self, ws, last_row: int, group_ranges: list[tuple[int, int]]):
        fill_group = PatternFill("solid", fgColor="EEF6F8")
        text_left = Alignment(horizontal="left", vertical="center")
        text_center = Alignment(horizontal="center", vertical="center")
        text_right = Alignment(horizontal="right", vertical="center")
        border = Border(
            left=Side(style="thin", color="D7E6EA"),
            right=Side(style="thin", color="D7E6EA"),
            top=Side(style="thin", color="D7E6EA"),
            bottom=Side(style="thin", color="D7E6EA"),
        )
        data_font = Font(size=10)
        highlight_font = Font(size=10, bold=True)

        for group_index, (start_row, end_row) in enumerate(group_ranges, start=1):
            use_fill = group_index % 2 == 1
            for row_idx in range(start_row, end_row + 1):
                for col_idx in range(1, 8):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if use_fill:
                        cell.fill = fill_group
                    cell.font = data_font
                    cell.border = border
                    if col_idx == 3:
                        cell.alignment = text_left
                    elif col_idx >= 4:
                        cell.alignment = text_right
                    else:
                        cell.alignment = text_center

                    if col_idx >= 4:
                        cell.number_format = "#,##0.00"

                ws.cell(row=row_idx, column=1).number_format = "@"
                ws.cell(row=row_idx, column=2).number_format = "@"

            ws.cell(row=start_row, column=1).font = highlight_font
            ws.cell(row=start_row, column=2).font = highlight_font
            ws.cell(row=start_row, column=7).font = highlight_font

            if end_row > start_row:
                for col_idx in (1, 2, 7):
                    ws.merge_cells(start_row=start_row, start_column=col_idx, end_row=end_row, end_column=col_idx)
                    ws.cell(row=start_row, column=col_idx).alignment = text_center

        widths = {
            1: 7,
            2: 12,
            3: 24,
            4: 18,
            5: 18,
            6: 16,
            7: 16,
        }
        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        if last_row >= 2:
            ws.freeze_panes = "A2"

    def _auto_style(self, ws, n_cols: int, last_row: int):
        fill_even = PatternFill("solid", fgColor="EEF6F8")
        text_left = Alignment(horizontal="left", vertical="center")
        text_center = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style="thin", color="D7E6EA"),
            right=Side(style="thin", color="D7E6EA"),
            top=Side(style="thin", color="D7E6EA"),
            bottom=Side(style="thin", color="D7E6EA"),
        )
        data_font = Font(size=10)

        for row_idx in range(2, last_row + 1):
            for col_idx in range(1, n_cols + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if row_idx % 2 == 0:
                    cell.fill = fill_even
                cell.font = data_font
                cell.border = border
                if col_idx == 3:
                    cell.alignment = text_left
                else:
                    cell.alignment = text_center

                if col_idx >= 4:
                    cell.number_format = "#,##0.00"

            ws.cell(row=row_idx, column=1).number_format = "@"
            ws.cell(row=row_idx, column=2).number_format = "@"

        for col_idx in range(1, n_cols + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(1, last_row + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                max_len = max(max_len, len(str(value or "")))

            if col_idx == 1:
                width = 12
            elif col_idx == 2:
                width = 14
            elif col_idx == 3:
                width = min(max(max_len + 2, 28), 42)
            else:
                width = min(max(max_len + 2, 14), 18)

            ws.column_dimensions[col_letter].width = width